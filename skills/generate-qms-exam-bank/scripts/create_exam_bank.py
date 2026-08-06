import argparse
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_HEADERS = [
    "题干",
    "类型",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "正确答案",
    "解析",
    "适用角色",
    "问题类别",
    "知识标题",
]


def normalize_row(row):
    return [row.get(header) for header in REQUIRED_HEADERS]


def validate_rows(rows, filename):
    if not rows:
        raise ValueError(f"{filename}: rows is empty")
    for index, row in enumerate(rows, start=2):
        qtype = row.get("类型")
        answer = row.get("正确答案")
        if not row.get("题干"):
            raise ValueError(f"{filename}: row {index} missing 题干")
        if qtype == "单选题":
            if answer not in {"A", "B", "C", "D"}:
                raise ValueError(f"{filename}: row {index} 单选题 answer must be A/B/C/D")
            for key in ["选项A", "选项B", "选项C", "选项D"]:
                if not row.get(key):
                    raise ValueError(f"{filename}: row {index} 单选题 missing {key}")
        elif qtype == "判断题":
            if row.get("选项A") != "正确" or row.get("选项B") != "错误":
                raise ValueError(f"{filename}: row {index} 判断题 options must be 正确/错误")
            if answer not in {"正确", "错误"}:
                raise ValueError(f"{filename}: row {index} 判断题 answer must be 正确 or 错误")
        else:
            raise ValueError(f"{filename}: row {index} unsupported 类型 {qtype!r}")
        for key in ["解析", "适用角色", "问题类别", "知识标题"]:
            if not row.get(key):
                raise ValueError(f"{filename}: row {index} missing {key}")


def copy_cell_style(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    dst.number_format = src.number_format
    dst.alignment = copy(src.alignment)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.font = copy(src.font)


def write_workbook(template, output_path, rows):
    wb = load_workbook(template)
    if "题库模板" not in wb.sheetnames:
        raise ValueError("template missing sheet: 题库模板")
    ws = wb["题库模板"]
    headers = [ws.cell(1, col).value for col in range(1, len(REQUIRED_HEADERS) + 1)]
    if headers != REQUIRED_HEADERS:
        raise ValueError(f"template headers mismatch: {headers}")

    style_row = 2 if ws.max_row >= 2 else 1
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row_index, row in enumerate(rows, start=2):
        values = normalize_row(row)
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col_index)
            copy_cell_style(ws.cell(style_row, col_index), cell)
            cell.value = value
            cell.alignment = copy(cell.alignment)
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="center")

    ws.freeze_panes = "A2"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Create QMS exam-bank workbooks from a JSON spec.")
    parser.add_argument("--spec", required=True, help="Path to JSON spec.")
    args = parser.parse_args()

    spec_path = Path(args.spec)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    template = Path(spec["template"])
    output_dir = Path(spec["output_dir"])

    created = []
    for workbook in spec["workbooks"]:
        filename = workbook["output_filename"]
        rows = workbook["rows"]
        validate_rows(rows, filename)
        output_path = output_dir / filename
        write_workbook(template, output_path, rows)
        created.append(str(output_path))

    for path in created:
        print(path)


if __name__ == "__main__":
    main()
